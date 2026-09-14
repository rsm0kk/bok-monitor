# -*- coding: utf-8 -*-
"""
extract.py — 한국은행 통화신용정책보고서 원본 데이터(xlsx) + 수출통계 104품목(xlsx) → out/data.json

  python extract.py                # data/ 안에서 가장 최근 파일을 자동으로 고른다
  python extract.py --bok=PATH --exp=PATH

원칙
  - 값이 없으면 null 로 둔다. 0 이나 추정치로 채우지 않는다.
  - 시트의 항목명·단위는 그대로 옮긴다(한글 라벨은 화면에서 그대로 쓴다).
  - 2026-09 (E) 같은 잠정 추정월은 flag 를 붙이고 증감률은 계산하지 않는다(원본 규약 그대로).
"""
import sys, os, re, json, glob, datetime as dt
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'data')
OUT = os.path.join(HERE, 'out')


def arg(name):
    for a in sys.argv[1:]:
        if a.startswith('--' + name + '='):
            return a.split('=', 1)[1]
    return None


def newest(pattern):
    files = sorted(glob.glob(os.path.join(DATA, pattern)), key=lambda f: (os.path.getmtime(f), os.path.basename(f)))
    if not files:
        raise SystemExit('파일 없음: ' + pattern)
    return files[-1]


def num(v):
    if v is None or v == '' or v == '–' or v == '-':
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    try:
        return round(float(str(v).replace(',', '')), 4)
    except ValueError:
        return None


def label(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    if v is None:
        return None
    return str(v).strip()


# ───────────────────────── 한국은행 시트 ─────────────────────────
def bok_sheet(wb, name):
    """'항목' 행을 찾아 열 이름을 읽고, 그 아래 단위 행을 건너뛴 뒤 데이터 행을 돌려준다."""
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = None
    for i, r in enumerate(rows[:12]):
        if r and r[0] is not None and str(r[0]).strip() in ('항목', '연%'):
            hdr_i = i
            break
    if hdr_i is None:
        raise SystemExit(name + ': 항목 행을 못 찾음')
    header = [label(c) for c in rows[hdr_i]]
    unit_row = rows[hdr_i + 1] if hdr_i + 1 < len(rows) else []
    units = [label(c) for c in unit_row]
    start = hdr_i + 1
    first = unit_row[0] if unit_row else None
    is_unit_row = bool(unit_row) and (first is None or not isinstance(first, (int, float, dt.datetime))) \
        and all(not isinstance(c, (int, float)) for c in unit_row[1:])
    if is_unit_row:
        start = hdr_i + 2
    else:
        units = []
    data = []
    for r in rows[start:]:
        if not r or r[0] is None or (isinstance(r[0], str) and not r[0].strip()):
            continue
        lab = label(r[0])
        vals = [num(c) for c in r[1:len(header)]]
        if all(v is None for v in vals):
            continue
        data.append([lab] + vals)
    names = [h for h in header[1:] if h]
    return {'title': label(rows[2][0]) if len(rows) > 2 else name,
            'source': label(rows[3][0]) if len(rows) > 3 else None,
            'cols': names, 'units': [u for u in units[1:len(header)]] if units else [], 'rows': data}


def series(sheet, keys=None):
    """rows → {dates:[...], <key>:[...]} 형태. keys 가 없으면 열 이름을 그대로 키로 쓴다."""
    cols = sheet['cols']
    keys = keys or cols
    out = {'dates': [r[0] for r in sheet['rows']], 'cols': cols, 'units': sheet['units'],
           'title': sheet['title'], 'source': sheet['source']}
    for j, k in enumerate(keys):
        out[k] = [r[j + 1] if j + 1 < len(r) else None for r in sheet['rows']]
    return out


def compress_steps(dates, vals):
    """일별 계단형 계열(기준금리)을 변경 시점만 남긴다."""
    out = []
    prev = object()
    for d, v in zip(dates, vals):
        if v != prev:
            out.append({'date': d, 'rate': v})
            prev = v
    return out


def extract_bok(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def S(n, k=None):
        return series(bok_sheet(wb, n), k)

    b = {}
    b['global'] = S('개요1', ['msci', 'dxy', 'ust10', 'jgb10'])
    b['gdpHalf'] = S('개요2', ['actual', 'forecast'])
    b['cpiOverview'] = S('개요3', ['cpi', 'cpiF', 'core', 'coreF'])
    b['domesticDaily'] = S('개요4', ['ktb3', 'corpAA', 'kospi', 'usdkrw'])
    b['housePrice'] = S('개요5', ['seoul', 'capital', 'noncapital'])
    br = S('개요7', ['rate'])
    b['baseRateSteps'] = compress_steps(br['dates'], br['rate'])
    b['baseRateLast'] = {'date': br['dates'][-1], 'rate': br['rate'][-1]}
    b['dailyExport'] = S('Ⅰ-11', ['exp', 'imp'])
    b['gdpQ'] = S('Ⅰ-15', ['qoq', 'qoqF', 'yoy', 'yoyF'])
    b['cpiMonthly'] = S('Ⅰ-16', ['mom', 'yoy'])
    b['coreMeasures'] = S('Ⅰ-19', ['exFoodEnergy', 'exAgriOil', 'exAdmin'])
    b['inflExp'] = S('Ⅰ-21', ['public1y', 'expert1y', 'consShort', 'consLong'])
    b['cpiForecastQ'] = S('Ⅰ-22', ['cpi', 'cpiF', 'core', 'coreF'])
    b['ktb'] = S('Ⅰ-23', ['y3', 'y10', 'spread'])
    b['credit'] = S('Ⅰ-26', ['aa', 'a', 'ktb3', 'aMinusAa', 'aMinusKtb', 'aaMinusKtb'])
    b['kospi'] = S('Ⅰ-28', ['kospi', 'vkospi'])
    b['foreignEq'] = S('Ⅰ-30', ['net', 'share'])
    b['houseTx'] = S('Ⅰ-33', ['priceAll', 'priceCap', 'txAll', 'txCap'])
    b['houseLoan'] = S('Ⅰ-34', ['bankMortgage', 'bankOther', 'nonbank', 'total'])
    b['fx'] = S('Ⅰ-35', ['usdkrw', 'neer'])
    b['growthContrib'] = S('Ⅱ-4', ['consumption', 'construction', 'facilities', 'exports', 'other', 'gdp'])
    b['itExport'] = S('Ⅱ-6', ['it', 'nonIt', 'total'])
    b['fcstGrowth'] = S('Ⅱ-10', ['may', 'aug'])
    b['fcstInfl'] = S('Ⅱ-11', ['cpiMay', 'cpiAug', 'coreMay', 'coreAug'])
    b['exportDomestic'] = S('Ⅱ-19', ['exports', 'exportsF', 'consumption', 'consumptionF'])
    b['inflPath'] = S('Ⅱ-20', ['cpi', 'cpiF', 'core', 'coreF'])
    b['houseCSI'] = S('Ⅱ-21', ['csi'])
    d1 = bok_sheet(wb, 'Ⅱ-22')
    d2 = bok_sheet(wb, 'Ⅱ-23')
    b['dots'] = {'may': [{'rate': num(r[0]), 'n': r[1]} for r in d1['rows']],
                 'aug': [{'rate': num(r[0]), 'n': r[1]} for r in d2['rows']]}
    b['tot'] = S('Ⅲ-3', ['tot', 'semi', 'otherIt', 'chem', 'metal'])
    b['itProfit'] = S('Ⅲ-6', ['itOp', 'itSales', 'otherOp', 'otherSales'])
    b['capex'] = S('참고 Ⅰ-1.1', ['amzn', 'msft', 'googl', 'meta', 'orcl', 'growth', 'annualGrowth'])
    b['semiExportQ'] = S('참고 Ⅰ-1.2', ['v'])
    b['aiRevenue'] = S('참고 Ⅰ-1.4', ['openai', 'anthropic'])
    b['dcGW'] = S('참고 Ⅰ-1.5', ['gw'])
    b['bondCapex'] = S('참고 Ⅰ-1.6', ['bond', 'capex', 'ratio'])
    b['semiContrib'] = S('참고 Ⅰ-4.3', ['samsung', 'hynix', 'total'])
    b['investorNet'] = S('참고 Ⅰ-4.6', ['foreign', 'retail', 'inst'])
    b['levEtf'] = S('참고 Ⅰ-4.8', ['index', 'sector', 'single'])
    b['singleEtf'] = S('참고 Ⅰ-4.10', ['hk', 'kr', 'us', 'uk'])
    b['policyRates'] = S('참고 Ⅱ-2.2', ['rate', 'chg'])
    b['fedExp'] = S('참고 Ⅱ-2.3', ['jan', 'apr', 'aug'])
    last = max(x['dates'][-1] for x in (b['ktb'], b['kospi'], b['fx'], b['credit']) if x['dates'])
    b['asOf'] = last
    return b


# ───────────────────────── 수출통계 시트 ─────────────────────────
CATS = ['반도체', 'IT 하드웨어', '이차전지', 'K-뷰티', '헬스케어', '전력 인프라', 'K-푸드',
        '기계·장비', '자동차·부품', '방산', '조선·해양', '석유·화학']


def extract_exports(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ex = {'categories': [], 'source': None}
    sm = [list(r) for r in wb['Summary'].iter_rows(values_only=True)]
    meta = {}
    for r in sm:
        if r and len(r) > 2 and r[1] is not None and isinstance(r[1], str):
            meta[str(r[1]).strip()] = r[2]
    ex['meta'] = {k: (str(v) if v is not None else None) for k, v in meta.items()}
    ex['official'] = []
    i0 = next((i for i, r in enumerate(sm) if r and len(r) > 1 and r[1] == '구분'), None)
    if i0 is not None:
        for r in sm[i0 + 1:]:
            if not r or r[1] is None:
                break
            ex['official'].append({'item': r[1], 'gov': r[2], 'ours': r[3], 'verdict': r[4], 'note': r[5]})
    months = None
    for cat in CATS:
        ws = wb[cat]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        names_row, hs_row = rows[4], rows[5]
        blocks = [(j, names_row[j]) for j in range(len(names_row)) if names_row[j] and j >= 2]
        data_rows = [r for r in rows[7:] if r and r[1] and re.match(r'^\d{4}-\d{2}', str(r[1]))]
        mlist = [str(r[1]) for r in data_rows]
        if months is None:
            months = mlist
        c = {'name': str(names_row[2]).replace('◆ ', '').replace(' 합계', ''), 'sheet': cat,
             'total': [num(r[2]) for r in data_rows], 'mm': [num(r[3]) for r in data_rows],
             'yoy': [num(r[4]) for r in data_rows], 'items': []}
        for j, nm in blocks[1:]:
            item = {'name': str(nm), 'hs': str(hs_row[j] or '').replace('HS ', ''),
                    'amt': [num(r[j]) for r in data_rows], 'mm': [num(r[j + 1]) for r in data_rows],
                    'yoy': [num(r[j + 2]) for r in data_rows], 'wt': [num(r[j + 3]) for r in data_rows],
                    'asp': [num(r[j + 4]) for r in data_rows]}
            c['items'].append(item)
        ex['categories'].append(c)
    ex['months'] = months
    ex['estimateMonths'] = [m for m in months if '(E)' in m]
    ex['latestActual'] = [m for m in months if '(E)' not in m][-1]
    mo = [list(r) for r in wb['모멘텀'].iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(mo) if r and len(r) > 1 and r[1] == '순위')
    ex['momentum'] = []
    for r in mo[hi + 1:]:
        if not r or r[1] is None:
            continue
        ex['momentum'].append({'rank': r[1], 'cat': r[2], 'item': r[3], 'hs': r[4], 'amt': num(r[5]),
                               'yoy': num(r[6]), 'baseFlag': r[7], 'prevYear': num(r[8]), 'accel': num(r[9]),
                               'avg3': num(r[10]), 'mm': num(r[11]), 'cum12': num(r[12]), 'cumYoy': num(r[13]),
                               'wtYoy': num(r[14]), 'aspYoy': num(r[15])})
    ex['momentumNote'] = str(mo[2][1]) if len(mo) > 2 and mo[2][1] else None
    es_name = next(s for s in wb.sheetnames if '근거' in s)
    es = [list(r) for r in wb[es_name].iter_rows(values_only=True)]
    est = {'sheet': es_name, 'meta': {}, 'rows': []}
    for r in es[:14]:
        if r and len(r) > 3 and r[2] and r[3] is not None and isinstance(r[2], str):
            est['meta'][r[2].strip()] = str(r[3])
    hi = next(i for i, r in enumerate(es) if r and len(r) > 1 and r[1] == '순위')
    hdr = es[hi]
    for r in es[hi + 1:]:
        if not r or r[1] is None:
            continue
        est['rows'].append({'rank': r[1], 'cat': r[2], 'item': r[3], 'hs': r[4], 'd10': num(r[5]),
                            'est': num(r[6]), 'prev': num(r[7]), 'capture': num(r[8]), 'prevYear': num(r[10]),
                            'wt': num(r[12]), 'asp': num(r[13]), 'note': r[14] if len(r) > 14 else None})
    est['header'] = [label(h) for h in hdr[1:15]]
    ex['estimate'] = est
    co = [list(r) for r in wb['국가별'].iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(co) if r and len(r) > 1 and r[1] == '카테고리')
    rows = [r for r in co[hi + 1:] if r and r[1]]
    ym = sorted({str(r[6]) for r in rows})
    keep = set(ym[-12:])
    country = {}
    for r in rows:
        if str(r[6]) not in keep:
            continue
        key = str(r[2])
        d = country.setdefault(key, {'cat': r[1], 'hs': r[3], 'countries': {}})
        cc = d['countries'].setdefault(str(r[5]), {'code': r[4], 'months': {}})
        cc['months'][str(r[6])] = {'amt': num(r[7]), 'yoy': num(r[8]), 'cum12': num(r[9]),
                                   'cumYoy': num(r[10]), 'share': num(r[11])}
    ex['country'] = {'months': ym[-12:], 'items': country}
    ex['source'] = ex['meta'].get('출처')
    ex['asOf'] = ex['meta'].get('작성 기준일')
    return ex


def main():
    bok_path = arg('bok') or newest('통화신용정책보고서 원본 데이터*.xlsx')
    exp_path = arg('exp') or newest('한국 수출통계*.xlsx')
    print('BOK :', os.path.basename(bok_path))
    print('EXP :', os.path.basename(exp_path))
    out = {'generatedAt': dt.datetime.now().strftime('%Y-%m-%d %H:%M'),
           'files': {'bok': os.path.basename(bok_path), 'exp': os.path.basename(exp_path)},
           'bok': extract_bok(bok_path), 'exports': extract_exports(exp_path)}
    os.makedirs(OUT, exist_ok=True)
    p = os.path.join(OUT, 'data.json')
    with open(p, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, separators=(',', ':'))
    print('→ %s · %.0fKB · BOK asOf %s · 수출 최신실측 %s' % (p, os.path.getsize(p) / 1024, out['bok']['asOf'], out['exports']['latestActual']))


if __name__ == '__main__':
    main()
